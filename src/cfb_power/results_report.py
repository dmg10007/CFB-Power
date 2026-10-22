"""Score the latest eligible saved pregame prediction for each game."""

from __future__ import annotations

import hashlib
import json
import math
import shutil
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class ReportError(ValueError):
    """Inputs cannot be safely used to publish a report."""


def stamp(value):
    try:
        if not isinstance(value, str):
            raise ValueError
        if value.endswith(" UTC"):
            return datetime.strptime(value, "%Y-%m-%d %H:%M UTC").replace(tzinfo=UTC)
        result = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if result.tzinfo is None:
            raise ValueError
        return result.astimezone(UTC)
    except (ValueError, TypeError):
        raise ReportError("A valid timezone-aware timestamp is required.") from None


def identifier(value):
    if type(value) is int and value > 0:
        return value
    if isinstance(value, str) and value.isdecimal() and int(value) > 0:
        return int(value)
    raise ReportError("Invalid game/team identifier.")


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def number(value):
    if type(value) is bool:
        raise ReportError("Boolean values are not valid predictions/scores.")
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ReportError("Non-numeric prediction/score.") from None
    if not math.isfinite(result):
        raise ReportError("Non-finite prediction/score.")
    return result


def forecasts(data_dir):
    rows, inputs = [], []
    for run in sorted(Path(data_dir).glob("runs/*/predictions_*")):
        if not run.is_dir():
            continue
        manifest_path = run / "manifest.json"
        try:
            before = digest(manifest_path)
            manifest = json.loads(manifest_path.read_text())
            if not isinstance(manifest, dict):
                raise ReportError("Forecast manifest must be an object.")
            if manifest.get("kind") != "upcoming_predictions" or manifest.get("status") != "ready":
                raise ReportError("Forecast run is not a ready upcoming-predictions run.")
            entries = [item for item in manifest["outputs"] if item["file"] == "predictions.csv"]
            if len(entries) != 1 or digest(run / "predictions.csv") != entries[0]["sha256"]:
                raise ReportError("Prediction CSV hash mismatch or missing hash.")
            csv_bytes = (run / "predictions.csv").read_bytes()
            if hashlib.sha256(csv_bytes).hexdigest() != entries[0]["sha256"]:
                raise ReportError("Prediction CSV changed while being read.")
            import io

            frame = pd.read_csv(io.BytesIO(csv_bytes), dtype=str, keep_default_na=False)
            required = {
                "game_id",
                "home_team_id",
                "away_team_id",
                "kickoff_utc",
                "low_history",
                "model",
                "home_expected_points",
                "away_expected_points",
                "projected_spread_home",
                "projected_total",
                "home_win_probability",
            }
            if (
                required - set(frame.columns)
                or type(manifest["counts"]["predicted"]) is not int
                or len(frame) != manifest["counts"]["predicted"]
            ):
                raise ReportError("Missing prediction columns or row-count mismatch.")
            created, as_of = stamp(manifest["created_at_utc"]), stamp(manifest["as_of_utc"])
            seen = set()
            for row in frame.to_dict("records"):
                if row["model"] != manifest["model"]:
                    raise ReportError("Prediction model disagrees with run manifest.")
                gid = identifier(row["game_id"])
                if gid in seen:
                    raise ReportError("Duplicate game in a prediction run.")
                seen.add(gid)
                row["game_id"] = gid
                for key in ("home_team_id", "away_team_id"):
                    row[key] = identifier(row[key])
                for key in (
                    "home_expected_points",
                    "away_expected_points",
                    "projected_spread_home",
                    "projected_total",
                    "home_win_probability",
                ):
                    row[key] = number(row[key])
                home, away = row["home_expected_points"], row["away_expected_points"]
                if min(home, away) < 0 or not 0 <= row["home_win_probability"] <= 1:
                    raise ReportError("Invalid predicted points/probability.")
                if not (
                    math.isclose(home - away, row["projected_spread_home"], abs_tol=1e-4)
                    and math.isclose(home + away, row["projected_total"], abs_tol=1e-4)
                ):
                    raise ReportError("Predicted score/margin/total disagree.")
                if row["low_history"] not in ("True", "False"):
                    raise ReportError("Invalid low-history flag.")
                row["low_history"] = row["low_history"] == "True"
                row.update(
                    created=created,
                    as_of=as_of,
                    saved_kickoff=stamp(row["kickoff_utc"]),
                    source_run=str(run.resolve()),
                )
                rows.append(row)
            if before != digest(manifest_path):
                raise ReportError("Forecast manifest changed during verification.")
            inputs.append(
                {
                    "source_run": str(run.resolve()),
                    "manifest_sha256": before,
                    "predictions_sha256": entries[0]["sha256"],
                    "created_at_utc": created.isoformat(),
                    "forecast_warnings": manifest.get("warnings", []),
                }
            )
        except (OSError, ValueError, KeyError, TypeError) as exc:
            raise ReportError(
                f"Cannot verify forecast run {run.name} ({type(exc).__name__}); report blocked."
            ) from None
    return rows, inputs


def compare(payload, candidates, season, week, season_type):
    if not isinstance(payload, list) or any(not isinstance(game, dict) for game in payload):
        raise ReportError("Results must be a list of game objects.")
    games, excluded, seen = [], [], set()
    for game in payload:
        gid = identifier(game.get("id"))
        if gid in seen:
            raise ReportError("Duplicate result game ID.")
        seen.add(gid)
        if type(game.get("season")) is not int or type(game.get("week")) is not int:
            raise ReportError("Invalid result season/week types.")
        if (game.get("season"), game.get("week"), game.get("seasonType")) != (
            season,
            week,
            season_type,
        ):
            raise ReportError("Result response is outside requested season/week/type.")
        classes = (game.get("homeClassification"), game.get("awayClassification"))
        if any(value not in ("fbs", "fcs") for value in classes):
            raise ReportError("Missing or unsupported team classification.")
        if "fcs" in classes:
            continue
        home_id, away_id = identifier(game.get("homeId")), identifier(game.get("awayId"))
        if home_id == away_id or type(game.get("completed")) is not bool:
            raise ReportError("Invalid result team identities/completed flag.")
        kickoff = (
            None
            if game.get("startTimeTBD") is True or not game.get("startDate")
            else stamp(game["startDate"])
        )
        actual_home = actual_away = None
        if game["completed"]:
            actual_home, actual_away = (
                number(game.get("homePoints")),
                number(game.get("awayPoints")),
            )
            if min(actual_home, actual_away) < 0 or not (
                actual_home.is_integer() and actual_away.is_integer()
            ):
                raise ReportError("Invalid final scores.")
        eligible = []
        for candidate in candidates:
            if candidate["game_id"] != gid:
                continue
            reason = None
            if (candidate["home_team_id"], candidate["away_team_id"]) != (home_id, away_id):
                reason = "team_identity_mismatch"
            elif kickoff is None:
                reason = "unconfirmed_kickoff"
            elif max(candidate["created"], candidate["as_of"]) >= min(
                kickoff, candidate["saved_kickoff"]
            ):
                reason = "not_created_and_asof_before_both_kickoffs"
            if reason:
                excluded.append(
                    {"game_id": gid, "source_run": candidate["source_run"], "reason": reason}
                )
            else:
                eligible.append(candidate)
        selected = (
            max(eligible, key=lambda row: (row["created"], row["source_run"])) if eligible else None
        )
        status = (
            "scored"
            if selected and game["completed"] and kickoff
            else (
                "unconfirmed_kickoff"
                if kickoff is None
                else "no_eligible_prediction"
                if not selected
                else "result_pending"
            )
        )
        row = {
            "game_id": gid,
            "home_team": game.get("homeTeam") or str(home_id),
            "away_team": game.get("awayTeam") or str(away_id),
            "home_team_id": home_id,
            "away_team_id": away_id,
            "kickoff_utc": kickoff.isoformat() if kickoff else "",
            "result_completed": game["completed"],
            "status": status,
            "actual_home_points": actual_home,
            "actual_away_points": actual_away,
            "home_expected_points": None,
            "away_expected_points": None,
            "projected_spread_home": None,
            "projected_total": None,
            "home_win_probability": None,
            "low_history": None,
            "model": "",
            "source_run": "",
            "forecast_created_at_utc": "",
            "forecast_as_of_utc": "",
            "actual_margin_home": None,
            "actual_total": None,
            "predicted_winner": "",
            "actual_winner": "",
            "winner_correct": None,
            "home_absolute_error": None,
            "away_absolute_error": None,
            "margin_absolute_error": None,
            "total_absolute_error": None,
            "brier": None,
        }
        if selected:
            for key in (
                "home_expected_points",
                "away_expected_points",
                "projected_spread_home",
                "projected_total",
                "home_win_probability",
                "low_history",
                "source_run",
            ):
                row[key] = selected[key]
            row["model"] = selected.get("model", "")
            row["forecast_created_at_utc"] = selected["created"].isoformat()
            row["forecast_as_of_utc"] = selected["as_of"].isoformat()
            margin = selected["projected_spread_home"]
            row["predicted_winner"] = "home" if margin > 0 else "away" if margin < 0 else "no_pick"
        if status == "scored":
            actual_margin = actual_home - actual_away
            row.update(
                actual_margin_home=actual_margin,
                actual_total=actual_home + actual_away,
                actual_winner="home"
                if actual_margin > 0
                else "away"
                if actual_margin < 0
                else "tie",
                home_absolute_error=abs(row["home_expected_points"] - actual_home),
                away_absolute_error=abs(row["away_expected_points"] - actual_away),
                margin_absolute_error=abs(row["projected_spread_home"] - actual_margin),
                total_absolute_error=abs(row["projected_total"] - actual_home - actual_away),
            )
            if actual_margin != 0:
                row["brier"] = (row["home_win_probability"] - int(actual_margin > 0)) ** 2
                if row["predicted_winner"] != "no_pick":
                    row["winner_correct"] = row["predicted_winner"] == row["actual_winner"]
        games.append(row)
    if not games:
        raise ReportError("No FBS-vs-FBS games in the requested scope.")
    return pd.DataFrame(games), pd.DataFrame(excluded, columns=["game_id", "source_run", "reason"])


def summarize(games):
    scored = games[games["status"] == "scored"]
    summaries = []
    for label, group in (
        ("all", scored),
        ("low_history", scored[scored["low_history"].eq(True)]),
        ("established_history", scored[scored["low_history"].eq(False)]),
    ):
        picks, probabilities = group["winner_correct"].dropna(), group["brier"].dropna()
        summaries.append(
            {
                "group": label,
                "games_scored": len(group),
                "winner_picks": len(picks),
                "correct_picks": int(picks.sum()) if len(picks) else 0,
                "winner_accuracy": float(picks.mean()) if len(picks) else None,
                "brier_games": len(probabilities),
                "brier_score": float(probabilities.mean()) if len(probabilities) else None,
                "score_mae": float(
                    (group["home_absolute_error"].mean() + group["away_absolute_error"].mean()) / 2
                )
                if len(group)
                else None,
                "margin_mae": float(group["margin_absolute_error"].mean()) if len(group) else None,
                "total_mae": float(group["total_absolute_error"].mean()) if len(group) else None,
            }
        )
    return pd.DataFrame(summaries)


def safe_text(value):
    return (
        "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
    )


def export_frame(frame):
    return frame.map(safe_text)


def workbook(path, summary, games, provenance):
    book = Workbook()
    book.remove(book.active)
    tables = {
        "Summary": summary,
        "Games": games,
        "Not Scored": games[games["status"] != "scored"],
        "Provenance": pd.DataFrame(provenance, columns=["item", "value"]),
    }
    for name, frame in tables.items():
        sheet = book.create_sheet(name)
        sheet.append(list(frame.columns))
        for values in export_frame(frame).itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in values])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123047")
        for index, column in enumerate(frame.columns, 1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = min(42, max(16, len(column) + 2))
            if column in ("winner_accuracy", "home_win_probability"):
                for cell in sheet[letter][1:]:
                    cell.number_format = "0.0%"
            if column.endswith("absolute_error") and len(frame):
                sheet.conditional_formatting.add(
                    f"{letter}2:{letter}{len(frame) + 1}",
                    ColorScaleRule(
                        start_type="min", start_color="C6EFCE", end_type="max", end_color="FFC7CE"
                    ),
                )
    book.save(path)


def generate_report(*, data_dir, season, week, client, season_type="regular"):
    if (
        type(season) is not int
        or season < 1
        or type(week) is not int
        or week < 0
        or season_type not in ("regular", "postseason")
    ):
        raise ReportError("Invalid season/week/type.")
    candidates, inputs = forecasts(data_dir)
    payload = client.get_json(
        "/games", {"year": season, "week": week, "seasonType": season_type, "classification": "fbs"}
    )
    captured = datetime.now(UTC).isoformat()
    games, exclusions = compare(payload, candidates, season, week, season_type)
    summary = summarize(games)
    report_id = datetime.now(UTC).strftime("%Y%m%dT%H%M%S%fZ") + "-" + uuid4().hex
    parent = Path(data_dir) / "reports"
    parent.mkdir(parents=True, exist_ok=True)
    staging, output = parent / ("." + report_id + ".staging"), parent / report_id
    staging.mkdir(exist_ok=False)
    counts = {
        "games_in_scope": len(games),
        "scored": int((games.status == "scored").sum()),
        "not_scored": int((games.status != "scored").sum()),
        "excluded_candidates": len(exclusions),
        "completed_without_prediction": int(
            (games.result_completed & (games.source_run == "")).sum()
        ),
    }
    manifest = {
        "kind": "weekly_results_report",
        "status": "ready",
        "schema_version": "results-report-v1",
        "created_at_utc": datetime.now(UTC).isoformat(),
        "results_captured_at_utc": captured,
        "scope": {"season": season, "week": week, "season_type": season_type},
        "selection": "latest_created_pregame_per_game; require created and asof before both saved and refreshed kickoff; tie-break by source_run path",
        "counts": counts,
        "partial": counts["not_scored"] > 0,
        "forecast_inputs": inputs,
        "notes": [
            "Final scores only; no predictions are generated.",
            "No sportsbook lines/ATS metrics.",
            "Source CSV hashes verified; source training datasets are not revalidated.",
        ],
    }
    try:
        (staging / "results.json").write_text(json.dumps(payload, indent=2, allow_nan=False) + "\n")
        for name, frame in (
            ("games.csv", games),
            ("not_scored.csv", games[games.status != "scored"]),
            ("summary.csv", summary),
            ("excluded_predictions.csv", exclusions),
        ):
            export_frame(frame).to_csv(staging / name, index=False)
        provenance = [
            ("Scope", json.dumps(manifest["scope"])),
            ("Results captured UTC", captured),
            ("Selection", manifest["selection"]),
            ("Counts", json.dumps(counts)),
            ("Partial report", str(manifest["partial"])),
        ]
        provenance.extend(("Forecast input", json.dumps(item)) for item in inputs)
        workbook(staging / "results_report.xlsx", summary, games, provenance)
        for item in inputs:
            run = Path(item["source_run"])
            if (
                digest(run / "manifest.json") != item["manifest_sha256"]
                or digest(run / "predictions.csv") != item["predictions_sha256"]
            ):
                raise ReportError("Forecast input changed before report publication.")
        manifest["outputs"] = [
            {"file": p.name, "sha256": digest(p)} for p in sorted(staging.iterdir())
        ]
        (staging / "manifest.json").write_text(
            json.dumps(manifest, indent=2, allow_nan=False) + "\n"
        )
        staging.rename(output)
        return output
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
