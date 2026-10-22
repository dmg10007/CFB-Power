from __future__ import annotations

from datetime import UTC, datetime
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Paths
from .team_names import enrich_predictions_frame, enrich_stats_frame, resolve_team_names


def _excel_safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Copy a frame and remove timezone metadata Excel cannot serialize."""
    safe = frame.copy()
    for column in safe.columns:
        if isinstance(safe[column].dtype, pd.DatetimeTZDtype):
            safe[column] = safe[column].dt.tz_localize(None)
    return safe


def _get_cell_style(column_name: str, value: Any) -> tuple[str | None, Alignment]:
    """Determine Excel number format and alignment for a column value."""
    col_lower = column_name.lower()

    # Identifiers and dates
    if any(k in col_lower for k in ("_id", "date", "kickoff", "season", "week")):
        return None, Alignment(horizontal="center", vertical="center")

    # String / Text
    if isinstance(value, str):
        return None, Alignment(horizontal="left", vertical="center")

    # Numeric formatting
    if isinstance(value, (int, float, np.integer, np.floating)) and not pd.isna(value):
        # Probabilities, percentages, and rates
        if any(
            k in col_lower
            for k in (
                "probability",
                "win_prob",
                "rate",
                "pct",
                "percent",
                "accuracy",
                "3rd_down",
                "red_zone",
            )
        ):
            return "0.0%", Alignment(horizontal="right", vertical="center")

        # High precision metrics
        if any(k in col_lower for k in ("brier", "log_loss")):
            return "0.000", Alignment(horizontal="right", vertical="center")

        # Spreads and score margins
        if any(k in col_lower for k in ("spread", "margin", "edge")):
            return "+0.0;-0.0;0.0", Alignment(horizontal="right", vertical="center")

        # Ratings, expected scores, averages, and continuous metrics
        if any(
            k in col_lower
            for k in (
                "rating",
                "srs",
                "sos",
                "form",
                "expected_points",
                "pts_avg",
                "interval",
                "mae",
            )
        ):
            return "0.0", Alignment(horizontal="right", vertical="center")

        # Counts, scores, and integers
        if isinstance(value, (int, np.integer)) or (
            isinstance(value, float) and value.is_integer()
        ):
            return "#,##0", Alignment(horizontal="right", vertical="center")

        return "0.00", Alignment(horizontal="right", vertical="center")

    return None, Alignment(horizontal="left", vertical="center")


def _write_dataframe(ws, frame: pd.DataFrame) -> None:
    safe_frame = _excel_safe_frame(frame)
    ws.append(list(safe_frame.columns))
    for row in safe_frame.itertuples(index=False, name=None):
        ws.append(list(row))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Styling constants
    header_fill = PatternFill("solid", fgColor="1B365D")  # Premium Classic Navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill("solid", fgColor="F8FAFC")  # Subtle slate/ice
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    border_side = Side(style="thin", color="E2E8F0")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Style header row
    ws.row_dimensions[1].height = 26
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    # Style data rows and cells
    for row_idx, row_cells in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        ws.row_dimensions[row_idx].height = 20
        fill = zebra_fill if row_idx % 2 == 0 else white_fill
        for col_idx, cell in enumerate(row_cells):
            col_name = str(safe_frame.columns[col_idx])
            val = cell.value
            num_fmt, align = _get_cell_style(col_name, val)
            cell.fill = fill
            cell.border = cell_border
            cell.alignment = align
            if num_fmt is not None:
                cell.number_format = num_fmt

    # Auto-fit column widths with padding
    for index, column in enumerate(safe_frame.columns, start=1):
        values = [str(column), *[str(value) for value in safe_frame[column].head(100).fillna("")]]
        max_len = max(len(value) for value in values)
        ws.column_dimensions[get_column_letter(index)].width = min(max(max_len + 3, 10), 32)


def export_all(
    predictions: pd.DataFrame,
    team_stats: pd.DataFrame,
    historical_results: pd.DataFrame,
    performance: pd.DataFrame,
    paths: Paths,
    power_rankings: pd.DataFrame | None = None,
    team_names: dict[str | int, str] | None = None,
) -> None:
    """Export predictions, team statistics, and evaluation metrics to CSV and formatted Excel.

    Parameters
    ----------
    predictions : pd.DataFrame
        Forecasted matchups with win probabilities and projected scores.
    team_stats : pd.DataFrame
        Team-level pregame rolling features.
    historical_results : pd.DataFrame
        Completed games with actual vs forecasted comparisons.
    performance : pd.DataFrame
        Model evaluation metrics (accuracy, MAE, log loss, Brier score).
    paths : Paths
        Directory paths configuration for export destination.
    power_rankings : pd.DataFrame | None, optional
        Opponent-adjusted team power rankings table.
    team_names : dict[str | int, str] | None, optional
        Mapping from team IDs to human-readable team names.
    """
    paths.ensure()

    # Resolve team names if not explicitly provided
    resolved_names = team_names or resolve_team_names(paths.data_dir)

    # Enrich DataFrames with readable team names and matchups
    enriched_predictions = enrich_predictions_frame(predictions, resolved_names)
    enriched_history = enrich_predictions_frame(historical_results, resolved_names)
    enriched_stats = enrich_stats_frame(team_stats, resolved_names)

    enriched_power: pd.DataFrame | None = None
    if power_rankings is not None:
        enriched_power = power_rankings.copy()
        if "team_name" not in enriched_power.columns and "team_id" in enriched_power.columns:
            enriched_power["team_name"] = (
                enriched_power["team_id"]
                .map(resolved_names)
                .fillna(enriched_power["team_id"].astype(str))
            )
            # Move team_name next to rank or team_id
            cols = list(enriched_power.columns)
            cols.remove("team_name")
            insert_pos = 2 if "rank" in cols else 1
            cols.insert(insert_pos, "team_name")
            enriched_power = enriched_power[cols]

    enriched_predictions.to_csv(paths.predictions_csv, index=False)
    enriched_stats.to_csv(paths.team_stats_csv, index=False)
    enriched_history.to_csv(paths.results_csv, index=False)
    performance.to_csv(paths.performance_csv, index=False)
    if enriched_power is not None:
        enriched_power.to_csv(paths.export_dir / "power_rankings.csv", index=False)

    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("README")
    readme.append(["CFB Power V1 Export"])
    readme.append(["Generated UTC", datetime.now(UTC).isoformat()])
    readme.append(
        ["Purpose", "Excel-ready output package; connect dashboard tables to these sheets or CSVs."]
    )
    readme.append(["Data rule", "All features are derived from games before the matchup date."])
    readme.append(["Model", "Two Ridge expected-score models with residual-based score intervals."])
    readme.append(
        [
            "Dashboard inputs",
            "Predictions, Team_Stats, Historical_Results, Model_Performance, Power_Rankings",
        ]
    )
    readme["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    readme["A1"].fill = PatternFill("solid", fgColor="1B365D")
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 100

    export_sheets: dict[str, pd.DataFrame] = {
        "Predictions": enriched_predictions,
        "Team_Stats": enriched_stats,
        "Historical_Results": enriched_history,
        "Model_Performance": performance,
    }
    if enriched_power is not None:
        export_sheets["Power_Rankings"] = enriched_power

    for sheet_name, frame in export_sheets.items():
        ws = wb.create_sheet(sheet_name)
        _write_dataframe(ws, frame)
        if sheet_name == "Predictions" and "home_win_probability" in frame.columns and len(frame):
            probability_col = get_column_letter(frame.columns.get_loc("home_win_probability") + 1)
            ws.conditional_formatting.add(
                f"{probability_col}2:{probability_col}{len(frame) + 1}",
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B",
                ),
            )
    wb.save(paths.workbook_xlsx)
