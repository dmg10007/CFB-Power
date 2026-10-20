"""Offline source audit and immutable experimental rating snapshots."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

import numpy as np
import pandas as pd

from .power_ratings import (
    RatingSettings, build_power_ratings, checked_games, eligible_history, utc_timestamp,
)

DATASET_SCHEMA = "dataset-basic-v1"
STAT_FIELDS = {"game_id", "team_id", "opponent_id", "is_home", "points", "points_allowed",
               "total_yards", "offensive_plays", "passing_yards", "rushing_yards", "turnovers",
               "takeaways", "third_down_attempts", "third_down_conversions", "red_zone_attempts",
               "red_zone_touchdowns", "penalties", "penalty_yards", "time_of_possession_seconds"}


def _digest(content):
    return hashlib.sha256(content).hexdigest()


def _validate_stat_links(games, stats):
    if not {"game_id", "team_id"} <= set(stats):
        raise ValueError("Stat keys are missing")
    if stats.duplicated(["game_id", "team_id"]).any():
        raise ValueError("Duplicate team-game statistics")
    completed = games[games.completed].set_index("game_id")
    if not set(stats.game_id) <= set(completed.index):
        raise ValueError("Statistics reference unknown or unfinished games")
    if len(stats) != 2 * len(completed):
        raise ValueError("Every completed game must have exactly two stat rows")
    for row in stats.itertuples(index=False):
        game = completed.loc[row.game_id]
        if row.team_id == game.home_team_id:
            opponent, own, other, home = game.away_team_id, game.home_score, game.away_score, True
        elif row.team_id == game.away_team_id:
            opponent, own, other, home = game.home_team_id, game.away_score, game.home_score, False
        else:
            raise ValueError("Statistics contain a team absent from its game")
        if (row.opponent_id != opponent or row.points != own or row.points_allowed != other
                or row.is_home not in (True, False) or bool(row.is_home) != home):
            raise ValueError("Statistics disagree with the schedule/final score")


def load_rating_dataset(directory):
    """Hash and parse the same bytes; do not import/train the legacy model stack."""
    directory = Path(directory)
    manifest_bytes = (directory / "manifest.json").read_bytes()
    manifest = json.loads(manifest_bytes)
    if (not isinstance(manifest, dict) or manifest.get("status") != "ready"
            or manifest.get("schema_version") != DATASET_SCHEMA):
        raise ValueError("Expected a ready dataset-basic-v1 manifest")
    entries = manifest.get("outputs")
    if (not isinstance(entries, list) or any(not isinstance(e, dict) for e in entries)
            or sorted(str(e.get("file")) for e in entries) != ["games.csv", "team_game_stats.csv"]):
        raise ValueError("Dataset must declare exactly games.csv and team_game_stats.csv")
    frames = {}
    for entry in entries:
        content = (directory / entry["file"]).read_bytes()
        if _digest(content) != entry.get("sha256"):
            raise ValueError(f"Hash mismatch: {entry['file']}")
        header = next(csv.reader(io.StringIO(content.decode("utf-8"))))
        if len(header) != len(set(header)):
            raise ValueError("Duplicate CSV column names")
        frame = pd.read_csv(io.BytesIO(content))
        if type(entry.get("rows")) is not int or len(frame) != entry["rows"]:
            raise ValueError(f"Row-count mismatch: {entry['file']}")
        frames[entry["file"]] = frame
    if manifest.get("offensive_plays_basis") != "passing_attempts_plus_rushing_attempts_proxy":
        raise ValueError("Unsupported or missing offensive-play basis")
    games = checked_games(frames["games.csv"])
    if "week" not in games:
        raise ValueError("Canonical week column is missing")
    stats = frames["team_game_stats.csv"]
    if not STAT_FIELDS <= set(stats):
        raise ValueError("Canonical team-stat columns are missing")
    _validate_stat_links(games, stats)
    if (directory / "manifest.json").read_bytes() != manifest_bytes:
        raise ValueError("Dataset manifest changed while loading")
    return manifest, games, stats, _digest(manifest_bytes)


def audit_source_coverage(games, stats, *, season, as_of):
    cutoff = utc_timestamp(as_of)
    checked = checked_games(games)
    if not {"game_id", "team_id"} <= set(stats):
        raise ValueError("Stat keys are missing")
    if stats.duplicated(["game_id", "team_id"]).any():
        raise ValueError("Duplicate team-game statistics")
    scope = checked[checked.season == season]
    history = eligible_history(scope, cutoff)
    expected = []
    for g in history.itertuples(index=False):
        expected.extend([(g.game_id, g.home_team_id), (g.game_id, g.away_team_id)])
    keys = pd.DataFrame(expected, columns=["game_id", "team_id"])
    selected = keys.merge(stats, on=["game_id", "team_id"], how="left", validate="one_to_one")
    rows = []
    def numeric(name):
        if name not in selected:
            return pd.Series(np.nan, index=selected.index, dtype=float)
        return pd.to_numeric(selected[name], errors="coerce").replace([np.inf, -np.inf], np.nan)
    def add(name, values, note, *, proxy=False):
        available = int(values.notna().sum())
        total = len(selected)
        status = "no_eligible_games" if not total else ("missing" if not available else
                 "partial" if available < total else "available")
        if proxy and available:
            status = "proxy" if available == total else "partial_proxy"
        rows.append({"feature": name, "eligible_team_games": total, "available_rows": available,
                     "unavailable_rows": total - available, "status": status, "note": note})
    def rate(name, numerator, denominator, note, *, proxy=False):
        num, den = numeric(numerator), numeric(denominator)
        valid = (den > 0) & num.notna()
        if name != "yards_per_offensive_play_proxy":
            valid &= num >= 0
        if name in ("third_down_conversion_rate", "red_zone_td_rate"):
            valid &= num <= den
        add(name, (num / den.where(den > 0)).where(valid), note, proxy=proxy)
    for name in ("points", "points_allowed", "turnovers", "takeaways", "penalties",
                 "penalty_yards", "time_of_possession_seconds"):
        add(name, numeric(name).where(numeric(name) >= 0), "Audited for coverage; only scores drive milestone-1 ratings")
    rate("yards_per_offensive_play_proxy", "total_yards", "offensive_plays",
         "Canonical denominator is attempts-plus-rushes proxy, not verified snaps", proxy=True)
    rate("third_down_conversion_rate", "third_down_conversions", "third_down_attempts",
         "Zero opportunities are undefined, not zero efficiency")
    rate("red_zone_td_rate", "red_zone_touchdowns", "red_zone_attempts",
         "Requires actual red-zone counts; no scoring-opportunity substitution")
    for name in ("passing_efficiency", "rushing_efficiency", "advanced_success_rate",
                 "roster_talent", "recruiting", "injuries", "starting_quarterback"):
        add(name, pd.Series(np.nan, index=selected.index),
            "Not integrated; provider access, historical coverage, and timestamps have not been verified")
    unresolved = scope[(scope.game_date < cutoff) & ~scope.completed]
    return {"season": season, "as_of_utc": cutoff.isoformat(),
            "coverage_denominator": "Completed, cutoff-eligible games represented in this dataset only",
            "counts": {"eligible_games": len(history), "eligible_team_games": len(selected),
                       "past_kickoff_unresolved": len(unresolved)},
            "external_api_requests": 0, "provider_access": "not_checked_offline",
            "result_availability_filter_present": "result_available_at_utc" in checked,
            "features": rows,
            "warnings": ["This audit does not establish completeness against the official FBS schedule",
                         "No API credentials, quota, or paid endpoint access are tested",
                         "Missing and zero-opportunity rates remain unavailable, never filled with zero"]}


def _json(path, value):
    path.write_text(json.dumps(value, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def _write_workbook(path, teams, audit):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    book.remove(book.active)
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return None if pd.isna(value) else value
    frames = [("Coverage", pd.DataFrame(audit["features"]))]
    if teams is not None:
        frames.insert(0, ("Team Ratings", teams))
    for name, frame in frames:
        sheet = book.create_sheet(name)
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            sheet.append([safe(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123047")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(len(str(c.value or "")) for c in column) + 2)
    book.save(path)


def run_rating_workflow(dataset_dir, data_dir, *, season, as_of, build=True, settings=None):
    directory = Path(dataset_dir)
    manifest, games, stats, manifest_hash = load_rating_dataset(directory)
    audit = audit_source_coverage(games, stats, season=season, as_of=as_of)
    snapshot = build_power_ratings(games, season=season, as_of=as_of, settings=settings) if build else None
    stamp = datetime.now(timezone.utc)
    identifier = stamp.strftime("%Y%m%dT%H%M%S%fZ") + "-" + uuid4().hex
    parent = Path(data_dir) / "rating_snapshots"
    destination, staging = parent / identifier, parent / ("." + identifier + ".staging")
    parent.mkdir(parents=True, exist_ok=True)
    staging.mkdir(exist_ok=False)
    try:
        _json(staging / "coverage.json", audit)
        pd.DataFrame(audit["features"]).to_csv(staging / "coverage.csv", index=False)
        if snapshot is not None:
            snapshot.teams.to_csv(staging / "team_ratings.csv", index=False)
        _write_workbook(staging / "power_ratings.xlsx", snapshot.teams if snapshot else None, audit)
        outputs = [{"file": p.name, "sha256": _digest(p.read_bytes())}
                   for p in sorted(staging.iterdir())]
        for entry in manifest["outputs"]:
            if _digest((directory / entry["file"]).read_bytes()) != entry["sha256"]:
                raise ValueError("Dataset artifact changed before publication")
        if _digest((directory / "manifest.json").read_bytes()) != manifest_hash:
            raise ValueError("Dataset manifest changed before publication")
        record = {"kind": "experimental_power_ratings" if build else "source_coverage_audit",
                  "status": "ready", "schema_version": "power-ratings-v1", "experimental": True,
                  "created_at_utc": stamp.isoformat(), "dataset_id": manifest.get("dataset_id"),
                  "dataset_manifest_sha256": manifest_hash, "source_outputs": manifest["outputs"],
                  "source_dataset_created_at_utc": manifest.get("created_at_utc"),
                  "source_dataset_path": str(directory.resolve()),
                  "season": season, "as_of_utc": utc_timestamp(as_of).isoformat(),
                  "outputs": outputs, "external_api_requests": 0,
                  "point_in_time_provenance_verified": False,
                  "notes": ["Existing datasets, forecasts, reports, and dashboard model are unchanged",
                            "Historical dataset revisions and source knowledge times are not reconstructed",
                            "Local snapshots are hashed; no fsync/crash-durability or external timestamp authority is claimed"],
                  "rating_model": snapshot.metadata if snapshot else None,
                  "base_points": snapshot.base_points if snapshot else None,
                  "home_field_points": snapshot.home_field_points if snapshot else None}
        _json(staging / "manifest.json", record)
        if destination.exists():
            raise FileExistsError("Snapshot already exists")
        os.rename(staging, destination)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return destination


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("audit", "build"))
    parser.add_argument("--dataset", type=Path, required=True)
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--season", type=int, required=True)
    parser.add_argument("--as-of", required=True, help="Explicit timezone-aware information cutoff")
    args = parser.parse_args(argv)
    try:
        output = run_rating_workflow(args.dataset, args.data_dir, season=args.season,
                                     as_of=args.as_of, build=args.command == "build")
    except (ValueError, OSError, KeyError, TypeError, StopIteration) as exc:
        print(f"Power-rating workflow failed: {exc}", file=sys.stderr)
        return 1
    print(f"Published experimental snapshot: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
