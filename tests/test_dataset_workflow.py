"""Dataset assembly and training tests using mocked CFBD snapshots; no API key needed."""

import copy
import hashlib
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from cfb_power.dataset_workflow import (
    DatasetError,
    assemble_dataset,
    discover_snapshots,
    main,
    train_from_dataset,
)
from cfb_power.ingestion.cfbd_batch import ingest_cfbd_batch

TEAMS = list(range(101, 117))


def make_game(game_id, week, home, away, completed=True, home_points=24, away_points=17):
    kickoff = datetime(2026, 9, 5, 16, tzinfo=timezone.utc) + timedelta(days=7 * (week - 1))
    return {
        "id": game_id, "season": 2026, "week": week, "seasonType": "regular",
        "startDate": kickoff.isoformat().replace("+00:00", "Z"), "startTimeTBD": False,
        "completed": completed, "neutralSite": False,
        "homeId": home, "homeTeam": f"Team {home}", "homeClassification": "fbs",
        "homePoints": home_points if completed else None,
        "awayId": away, "awayTeam": f"Team {away}", "awayClassification": "fbs",
        "awayPoints": away_points if completed else None,
    }


def make_box(game, rng):
    teams = []
    for side in ("home", "away"):
        passing, rushing = int(rng.integers(120, 300)), int(rng.integers(50, 200))
        attempts = int(rng.integers(20, 40))
        stats = {
            "thirdDownEff": f"{int(rng.integers(2, 8))}-{int(rng.integers(9, 16))}",
            "completionAttempts": f"{int(attempts * 0.6)}-{attempts}",
            "rushingAttempts": str(int(rng.integers(20, 45))),
            "totalYards": str(passing + rushing), "netPassingYards": str(passing),
            "rushingYards": str(rushing),
            "totalPenaltiesYards": f"{int(rng.integers(2, 10))}-{int(rng.integers(15, 90))}",
            "turnovers": str(int(rng.integers(0, 4))), "possessionTime": "30:00",
        }
        teams.append({
            "teamId": game[f"{side}Id"], "team": game[f"{side}Team"], "homeAway": side,
            "points": game[f"{side}Points"],
            "stats": [{"category": name, "stat": value} for name, value in stats.items()],
        })
    return {"id": game["id"], "teams": teams}


def week_games(rng, week, completed=True):
    order = [int(team) for team in rng.permutation(TEAMS)]
    games = []
    for index in range(0, len(order), 2):
        game_id = 9000 + week * 100 + index // 2
        points = (int(rng.integers(10, 45)), int(rng.integers(10, 45)))
        games.append(make_game(game_id, week, order[index], order[index + 1], completed, *points))
    return games


def capture(tmp_path, games, week):
    rng = np.random.default_rng(week + 1000)
    client = Mock()
    boxes = [[make_box(game, rng)] for game in games if game["completed"]]
    client.get_json.side_effect = [games, *boxes]
    return ingest_cfbd_batch(season=2026, season_type="regular", week=week, data_dir=tmp_path, client=client)


@pytest.fixture()
def world(tmp_path):
    rng = np.random.default_rng(7)
    weeks = {week: week_games(rng, week) for week in (1, 2, 3)}
    snapshots = [capture(tmp_path, weeks[week], week) for week in (1, 2, 3)]
    return tmp_path, weeks, snapshots


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_assemble_verifies_merges_and_preserves_inputs(world):
    tmp_path, _, snapshots = world
    active = tmp_path / "processed"
    (active / "games.csv").write_text("existing games", encoding="utf-8")
    (active / "team_game_stats.csv").write_text("existing stats", encoding="utf-8")
    before = {path: digest(path) for snapshot in snapshots for path in snapshot.iterdir()}
    dataset = assemble_dataset(snapshots, tmp_path)
    manifest = json.loads((dataset / "manifest.json").read_text())
    assert manifest["status"] == "ready"
    assert manifest["counts"] == {
        "snapshots": 3, "games": 24, "completed_games": 24, "upcoming_games": 0,
        "team_stat_rows": 48, "upcoming_replaced_by_completed": 0,
    }
    assert [item["run_id"] for item in manifest["inputs"]] == [path.name for path in snapshots]
    for artifact in manifest["outputs"]:
        assert digest(dataset / artifact["file"]) == artifact["sha256"]
    games = pd.read_csv(dataset / "games.csv")
    stats = pd.read_csv(dataset / "team_game_stats.csv")
    assert len(games) == 24 and len(stats) == 48
    assert stats["red_zone_attempts"].isna().all()
    assert {path: digest(path) for snapshot in snapshots for path in snapshot.iterdir()} == before
    assert (active / "games.csv").read_text() == "existing games"
    assert (active / "team_game_stats.csv").read_text() == "existing stats"
    assert not list((tmp_path / "datasets").glob(".*.staging"))


def test_identical_recaptures_are_deduplicated(tmp_path):
    games = week_games(np.random.default_rng(3), 1)
    first, second = capture(tmp_path, games, 1), capture(tmp_path, copy.deepcopy(games), 1)
    manifest = json.loads((assemble_dataset([first, second], tmp_path) / "manifest.json").read_text())
    assert manifest["counts"]["games"] == 8
    assert manifest["counts"]["team_stat_rows"] == 16


def test_completed_game_replaces_earlier_upcoming_record(tmp_path):
    rng = np.random.default_rng(5)
    completed = week_games(rng, 1)
    upcoming = copy.deepcopy(completed)
    for game in upcoming:
        game.update(completed=False, homePoints=None, awayPoints=None)
    first, second = capture(tmp_path, upcoming, 1), capture(tmp_path, completed, 1)
    manifest = json.loads((assemble_dataset([first, second], tmp_path) / "manifest.json").read_text())
    assert manifest["counts"]["completed_games"] == 8
    assert manifest["counts"]["upcoming_games"] == 0
    assert manifest["counts"]["upcoming_replaced_by_completed"] == 8


def test_conflicting_games_fail_without_publishing(tmp_path):
    games = week_games(np.random.default_rng(9), 1)
    changed = copy.deepcopy(games)
    changed[0]["homePoints"] += 1
    first, second = capture(tmp_path, games, 1), capture(tmp_path, changed, 1)
    with pytest.raises(DatasetError, match="conflicting"):
        assemble_dataset([first, second], tmp_path)
    assert not (tmp_path / "datasets").exists()


@pytest.mark.parametrize("case", ["tampered", "unready", "missing_manifest", "wrong_schema"])
def test_untrusted_snapshots_are_rejected(world, case):
    tmp_path, _, snapshots = world
    target = snapshots[0]
    manifest_path = target / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    if case == "tampered":
        with (target / "games.csv").open("a", encoding="utf-8") as handle:
            handle.write("\n")
    elif case == "unready":
        manifest["status"] = "failed"
        manifest_path.write_text(json.dumps(manifest))
    elif case == "wrong_schema":
        manifest["schema_version"] = "other"
        manifest_path.write_text(json.dumps(manifest))
    else:
        manifest_path.unlink()
    with pytest.raises(DatasetError):
        assemble_dataset(snapshots, tmp_path)
    assert not (tmp_path / "datasets").exists()


def test_publication_failure_removes_staging(world):
    tmp_path, _, snapshots = world
    with patch("cfb_power.dataset_workflow.Path.rename", side_effect=OSError("publication error")):
        with pytest.raises(DatasetError):
            assemble_dataset(snapshots, tmp_path)
    assert list((tmp_path / "datasets").iterdir()) == []


def test_training_uses_explicit_dataset_and_writes_only_to_new_run(world):
    tmp_path, _, snapshots = world
    active = tmp_path / "processed"
    (active / "games.csv").write_text("existing games", encoding="utf-8")
    dataset = assemble_dataset(snapshots, tmp_path)
    before = {path: digest(path) for path in dataset.iterdir()}
    run = train_from_dataset(dataset, tmp_path)
    run_manifest = json.loads((run / "run_manifest.json").read_text())
    assert run_manifest["evaluation_scope"] == "in_sample_pipeline_check"
    assert run_manifest["completed_games"] == 24
    assert (run / "models" / "score_models.joblib").exists()
    predictions = pd.read_csv(run / "exports" / "predictions.csv")
    assert len(predictions) == 24
    assert predictions["home_win_probability"].between(0, 1).all()
    assert {"README", "Predictions", "Team_Stats", "Historical_Results", "Model_Performance"} == set(
        load_workbook(run / "exports" / "cfb_forecast_export.xlsx", read_only=True).sheetnames
    )
    assert {path: digest(path) for path in dataset.iterdir()} == before
    assert (active / "games.csv").read_text() == "existing games"


def test_training_requires_multiple_weeks_before_creating_a_run(world):
    tmp_path, _, snapshots = world
    dataset = assemble_dataset(snapshots[:1], tmp_path)
    with pytest.raises(DatasetError, match="at least two"):
        train_from_dataset(dataset, tmp_path)
    assert not (tmp_path / "runs").exists()


def test_training_rejects_modified_dataset(world):
    tmp_path, _, snapshots = world
    dataset = assemble_dataset(snapshots, tmp_path)
    with (dataset / "games.csv").open("a", encoding="utf-8") as handle:
        handle.write("\n")
    with pytest.raises(DatasetError, match="hash mismatch"):
        train_from_dataset(dataset, tmp_path)


def test_cli_discovers_snapshots_assembles_and_trains(world, capsys):
    tmp_path, _, snapshots = world
    assert discover_snapshots(tmp_path) == sorted(snapshots)
    assert main(["--data-dir", str(tmp_path), "assemble", "--all"]) == 0
    dataset = next((tmp_path / "datasets").iterdir())
    assert "Original snapshots and active model inputs were not changed" in capsys.readouterr().out
    assert main(["--data-dir", str(tmp_path), "train", "--dataset", str(dataset)]) == 0
    assert "not out-of-sample accuracy" in capsys.readouterr().out


def test_cli_reports_missing_snapshots(tmp_path, capsys):
    assert main(["--data-dir", str(tmp_path), "assemble", "--all"]) == 1
    assert "at least one snapshot" in capsys.readouterr().err
