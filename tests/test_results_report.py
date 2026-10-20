"""Weekly report tests using local synthetic forecasts and mocked final scores."""
import json
from unittest.mock import Mock

import pandas as pd
import pytest
from openpyxl import load_workbook
from cfb_power import results_report as report


def game(gid=1, completed=True, home=31, away=24):
    return {"id": gid, "season": 2026, "week": 7, "seasonType": "regular",
            "homeId": gid * 10 + 1, "awayId": gid * 10 + 2, "homeTeam": "Home", "awayTeam": "Away",
            "homePoints": home, "awayPoints": away, "completed": completed,
            "homeClassification": "fbs", "awayClassification": "fbs", "startTimeTBD": False,
            "startDate": "2026-10-17T16:00:00Z" if gid == 1 else "2026-10-17T20:00:00Z"}


def forecast(tmp_path, name, ids=(1,), created="2026-10-16T01:00:00Z", asof=None, home=30, away=20, probability=0.7, low=False):
    run = tmp_path / "runs/dataset" / ("predictions_" + name)
    run.mkdir(parents=True)
    rows = [{"game_id": gid, "home_team_id": gid*10+1, "away_team_id": gid*10+2,
             "kickoff_utc": "2026-10-17 16:00 UTC" if gid == 1 else "2026-10-17 20:00 UTC",
             "home_expected_points": home, "away_expected_points": away, "projected_spread_home": home-away,
             "projected_total": home+away, "home_win_probability": probability, "low_history": low,
             "model": "rolling_average"} for gid in ids]
    pd.DataFrame(rows).to_csv(run / "predictions.csv", index=False)
    manifest = {"kind": "upcoming_predictions", "status": "ready", "created_at_utc": created,
                "as_of_utc": asof or created, "model": "rolling_average", "counts": {"predicted": len(rows)},
                "outputs": [{"file": "predictions.csv", "sha256": report.digest(run / "predictions.csv")}]}
    (run / "manifest.json").write_text(json.dumps(manifest))
    return run


def compare(tmp_path, games):
    candidates, _ = report.forecasts(tmp_path)
    return report.compare(games, candidates, 2026, 7, "regular")


def test_latest_is_selected_per_game_not_global_run(tmp_path):
    old = forecast(tmp_path, "old", ids=(1, 2))
    new = forecast(tmp_path, "new", created="2026-10-17T15:00:00Z", home=28, away=21, probability=0.6)
    forecast(tmp_path, "after", created="2026-10-18T01:00:00Z", asof="2026-10-16T01:00:00Z")
    games, excluded = compare(tmp_path, [game(), game(2, home=17, away=28)])
    assert games.source_run.tolist() == [str(new.resolve()), str(old.resolve())]
    assert len(games) == 2 and len(excluded) == 1
    summary = report.summarize(games).set_index("group").loc["all"]
    assert summary.winner_accuracy == 0.5
    assert summary.score_mae == pytest.approx(6.75)
    assert summary.margin_mae == pytest.approx(10.5)
    assert summary.total_mae == pytest.approx(5.5)
    assert summary.brier_score == pytest.approx(0.325)


@pytest.mark.parametrize("created,asof", [("2026-10-17T16:00:00Z", "2026-10-16T00:00:00Z"),
                                         ("2026-10-16T00:00:00Z", "2026-10-17T16:00:00Z")])
def test_actual_creation_and_asof_must_both_be_strictly_pregame(tmp_path, created, asof):
    forecast(tmp_path, "late", created=created, asof=asof)
    games, excluded = compare(tmp_path, [game()])
    assert games.iloc[0].status == "no_eligible_prediction" and len(excluded) == 1


def test_changed_kickoff_does_not_make_post_original_forecast_eligible(tmp_path):
    forecast(tmp_path, "late", created="2026-10-17T17:00:00Z")
    result = game()
    result["startDate"] = "2026-10-17T18:00:00Z"
    assert compare(tmp_path, [result])[0].iloc[0].status == "no_eligible_prediction"


def test_pending_and_missing_forecasts_remain_visible(tmp_path):
    forecast(tmp_path, "old")
    games, _ = compare(tmp_path, [game(completed=False, home=None, away=None), game(2)])
    assert games.status.tolist() == ["result_pending", "no_eligible_prediction"]
    assert report.summarize(games).iloc[0].games_scored == 0


def test_projected_tie_is_no_pick_but_has_brier(tmp_path):
    forecast(tmp_path, "tie", home=20, away=20, probability=0.5)
    games, _ = compare(tmp_path, [game()])
    assert games.iloc[0].predicted_winner == "no_pick"
    summary = report.summarize(games).iloc[0]
    assert summary.winner_picks == 0 and summary.brier_games == 1
    assert summary.brier_score == 0.25


def test_actual_tie_still_scores_errors_not_winner_or_brier(tmp_path):
    forecast(tmp_path, "old")
    games, _ = compare(tmp_path, [game(home=24, away=24)])
    summary = report.summarize(games).iloc[0]
    assert summary.games_scored == 1 and summary.winner_picks == 0 and summary.brier_games == 0
    assert summary.margin_mae == 10


def test_team_mismatch_and_unconfirmed_time_are_not_scored(tmp_path):
    forecast(tmp_path, "old", ids=(1, 2))
    first, second = game(), game(2)
    first["homeId"] = 999
    second["startTimeTBD"] = True
    games, exclusions = compare(tmp_path, [first, second])
    assert games.status.tolist() == ["no_eligible_prediction", "unconfirmed_kickoff"]
    assert set(exclusions.reason) == {"team_identity_mismatch", "unconfirmed_kickoff"}


def test_low_history_summary_is_separate(tmp_path):
    forecast(tmp_path, "low", low=True)
    games, _ = compare(tmp_path, [game()])
    summary = report.summarize(games).set_index("group")
    assert summary.loc["low_history", "games_scored"] == 1
    assert summary.loc["established_history", "games_scored"] == 0


def test_corrupt_forecast_blocks_before_any_api_call(tmp_path):
    run = forecast(tmp_path, "old")
    (run / "predictions.csv").write_text("tampered")
    client = Mock()
    with pytest.raises(report.ReportError, match="verify"):
        report.generate_report(data_dir=tmp_path, season=2026, week=7, client=client)
    client.get_json.assert_not_called()


def test_missing_manifest_does_not_silently_fall_back(tmp_path):
    forecast(tmp_path, "old")
    (tmp_path / "runs/dataset/predictions_broken").mkdir()
    with pytest.raises(report.ReportError):
        report.forecasts(tmp_path)


@pytest.mark.parametrize("field,value", [("created_at_utc", "2026-10-16 01:00:00"), ("as_of_utc", "not-a-date")])
def test_invalid_timestamp_blocks_report(tmp_path, field, value):
    run = forecast(tmp_path, "old")
    path = run / "manifest.json"
    manifest = json.loads(path.read_text())
    manifest[field] = value
    path.write_text(json.dumps(manifest))
    with pytest.raises(report.ReportError):
        report.forecasts(tmp_path)


def test_duplicate_result_rejected(tmp_path):
    with pytest.raises(report.ReportError, match="Duplicate"):
        compare(tmp_path, [game(), game()])


def test_deterministic_tiebreak(tmp_path):
    forecast(tmp_path, "a")
    latest = forecast(tmp_path, "z", home=20, away=10)
    assert compare(tmp_path, [game()])[0].iloc[0].source_run == str(latest.resolve())


def test_report_exports_hashes_and_does_not_modify_forecasts(tmp_path):
    run = forecast(tmp_path, "old")
    before = {p: p.read_bytes() for p in run.iterdir()}
    client = Mock()
    client.get_json.return_value = [game(), game(2)]
    output = report.generate_report(data_dir=tmp_path, season=2026, week=7, client=client)
    client.get_json.assert_called_once_with("/games", {"year": 2026, "week": 7, "seasonType": "regular", "classification": "fbs"})
    manifest = json.loads((output / "manifest.json").read_text())
    assert manifest["counts"]["scored"] == 1 and manifest["counts"]["not_scored"] == 1
    assert manifest["partial"] is True and manifest["forecast_inputs"][0]["source_run"] == str(run.resolve())
    for item in manifest["outputs"]:
        assert report.digest(output / item["file"]) == item["sha256"]
    assert load_workbook(output / "results_report.xlsx").sheetnames == ["Summary", "Games", "Not Scored", "Provenance"]
    assert all(p.read_bytes() == value for p, value in before.items())
    assert not list((tmp_path / "reports").glob(".*.staging"))


def test_no_saved_predictions_still_reports_coverage_gap(tmp_path):
    client = Mock()
    client.get_json.return_value = [game()]
    output = report.generate_report(data_dir=tmp_path, season=2026, week=7, client=client)
    manifest = json.loads((output / "manifest.json").read_text())
    assert manifest["counts"]["scored"] == 0 and manifest["counts"]["completed_without_prediction"] == 1


def test_formula_like_team_names_are_escaped(tmp_path):
    result = game()
    result["homeTeam"] = "=HYPERLINK(unsafe)"
    client = Mock()
    client.get_json.return_value = [result]
    output = report.generate_report(data_dir=tmp_path, season=2026, week=7, client=client)
    sheet = load_workbook(output / "results_report.xlsx")["Games"]
    assert sheet["B2"].data_type != "f" and sheet["B2"].value.startswith("'=")
    assert "'=HYPERLINK" in (output / "games.csv").read_text()


def test_export_failure_leaves_no_ready_report(tmp_path, monkeypatch):
    client = Mock()
    client.get_json.return_value = [game()]
    monkeypatch.setattr(report, "workbook", Mock(side_effect=OSError("disk failure")))
    with pytest.raises(OSError):
        report.generate_report(data_dir=tmp_path, season=2026, week=7, client=client)
    assert not list((tmp_path / "reports").iterdir())


@pytest.mark.parametrize("column,value", [("home_win_probability", "NaN"), ("projected_total", "999"), ("low_history", "unknown")])
def test_invalid_prediction_values_are_rejected_even_with_updated_hash(tmp_path, column, value):
    run = forecast(tmp_path, "old")
    frame = pd.read_csv(run / "predictions.csv", dtype=str)
    frame.loc[0, column] = value
    frame.to_csv(run / "predictions.csv", index=False)
    path = run / "manifest.json"
    manifest = json.loads(path.read_text())
    manifest["outputs"][0]["sha256"] = report.digest(run / "predictions.csv")
    path.write_text(json.dumps(manifest))
    with pytest.raises(report.ReportError):
        report.forecasts(tmp_path)


def test_inputs_changed_during_export_block_publication(tmp_path, monkeypatch):
    run = forecast(tmp_path, "old")
    client = Mock()
    client.get_json.return_value = [game()]
    def mutate(*args):
        (run / "predictions.csv").write_text("changed")
    monkeypatch.setattr(report, "workbook", mutate)
    with pytest.raises(report.ReportError, match="changed"):
        report.generate_report(data_dir=tmp_path, season=2026, week=7, client=client)
    assert not list((tmp_path / "reports").iterdir())
