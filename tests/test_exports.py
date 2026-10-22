import pandas as pd
from openpyxl import load_workbook

from cfb_power.config import Paths
from cfb_power.exports import export_all


def test_export_all_writes_timezone_aware_dates_to_excel(tmp_path):
    paths = Paths(
        root=tmp_path,
        data_dir=tmp_path / "data",
        export_dir=tmp_path / "data" / "exports",
        model_dir=tmp_path / "models",
    )
    predictions = pd.DataFrame(
        {
            "game_id": ["g1"],
            "game_date": pd.to_datetime(["2025-09-01T18:00:00Z"], utc=True),
            "home_win_probability": [0.6],
        }
    )
    team_stats = pd.DataFrame(
        {
            "game_id": ["g1"],
            "game_date": pd.to_datetime(["2025-09-01T18:00:00Z"], utc=True),
        }
    )
    export_all(predictions, team_stats, predictions.copy(), pd.DataFrame(), paths)
    assert paths.workbook_xlsx.exists()
    workbook = load_workbook(paths.workbook_xlsx, read_only=True)
    assert workbook.sheetnames == [
        "README",
        "Predictions",
        "Team_Stats",
        "Historical_Results",
        "Model_Performance",
    ]


def test_export_all_enriches_team_names_and_formatting(tmp_path):
    paths = Paths(
        root=tmp_path,
        data_dir=tmp_path / "data",
        export_dir=tmp_path / "data" / "exports",
        model_dir=tmp_path / "models",
    )
    predictions = pd.DataFrame(
        {
            "game_id": [101],
            "season": [2026],
            "week": [1],
            "game_date": pd.to_datetime(["2026-09-01T18:00:00Z"], utc=True),
            "home_team_id": [10],
            "away_team_id": [20],
            "neutral_site": [False],
            "projected_spread_home": [7.5],
            "home_win_probability": [0.725],
        }
    )
    team_stats = pd.DataFrame(
        {
            "game_id": [101],
            "game_date": pd.to_datetime(["2026-09-01T18:00:00Z"], utc=True),
            "team_id": [10],
            "opponent_id": [20],
            "points": [35],
        }
    )
    names = {10: "Ohio State", 20: "Michigan"}
    power = pd.DataFrame(
        {
            "rank": [1, 2],
            "team_id": [10, 20],
            "composite_rating": [25.4, 21.2],
        }
    )
    export_all(
        predictions,
        team_stats,
        predictions.copy(),
        pd.DataFrame([{"winner_accuracy": 0.75}]),
        paths,
        power_rankings=power,
        team_names=names,
    )

    # Verify CSV files have team names and matchups
    pred_csv = pd.read_csv(paths.predictions_csv)
    assert "home_team" in pred_csv.columns
    assert "away_team" in pred_csv.columns
    assert "matchup" in pred_csv.columns
    assert pred_csv["matchup"].iloc[0] == "Michigan @ Ohio State"
    assert pred_csv["home_team"].iloc[0] == "Ohio State"

    stats_csv = pd.read_csv(paths.team_stats_csv)
    assert "team_name" in stats_csv.columns
    assert stats_csv["team_name"].iloc[0] == "Ohio State"
    assert stats_csv["opponent_name"].iloc[0] == "Michigan"

    # Verify workbook sheets and cell formatting
    wb = load_workbook(paths.workbook_xlsx)
    assert "Power_Rankings" in wb.sheetnames
    pred_ws = wb["Predictions"]
    headers = [cell.value for cell in pred_ws[1]]
    assert "matchup" in headers
    assert "home_team" in headers
    # Check that home_win_probability has percentage formatting
    prob_idx = headers.index("home_win_probability") + 1
    assert pred_ws.cell(row=2, column=prob_idx).number_format == "0.0%"
