"""Tests for the warning about games past kickoff that have no recorded result."""

import json

import pandas as pd
from openpyxl import load_workbook
from test_upcoming import make_tables, make_upcoming, week_kickoff, write_dataset

from cfb_power.upcoming import predict_upcoming


def split_week_six(games, stats):
    """Week 6 upcoming, with four games kicking off four hours before the other four."""
    up_games, up_stats = make_upcoming(games, stats)
    ids = up_games.loc[up_games["week"] == 6, "game_id"].tolist()
    kickoff = week_kickoff(games, 6)
    up_games.loc[up_games["game_id"].isin(ids[:4]), "game_date"] = kickoff.isoformat()
    up_games.loc[up_games["game_id"].isin(ids[4:]), "game_date"] = (
        kickoff + pd.Timedelta(hours=4)
    ).isoformat()
    return up_games, up_stats, kickoff


def test_games_past_kickoff_without_results_trigger_a_warning(tmp_path):
    games, stats = make_tables()
    up_games, up_stats, kickoff = split_week_six(games, stats)
    run = predict_upcoming(
        write_dataset(tmp_path, up_games, up_stats), tmp_path, as_of=kickoff + pd.Timedelta(hours=1)
    )
    manifest = json.loads((run / "manifest.json").read_text())
    assert manifest["counts"]["past_kickoff_unresolved"] == 4
    assert any(
        "4 games are past kickoff with no recorded result" in warning
        for warning in manifest["warnings"]
    )
    readme = load_workbook(run / "upcoming_predictions.xlsx")["README"]
    warnings_cell = next(
        row[1] for row in readme.iter_rows(values_only=True) if row[0] == "Warnings"
    )
    assert "4 games are past kickoff" in warnings_cell


def test_no_unresolved_warning_when_every_game_is_still_upcoming(tmp_path):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    run = predict_upcoming(
        write_dataset(tmp_path, up_games, up_stats),
        tmp_path,
        as_of=week_kickoff(games, 6) - pd.Timedelta(days=1),
    )
    manifest = json.loads((run / "manifest.json").read_text())
    assert manifest["counts"]["past_kickoff_unresolved"] == 0
    assert not any("past kickoff" in warning for warning in manifest["warnings"])


def test_unrecaptured_week_is_flagged_even_when_the_data_age_looks_fine(tmp_path):
    games, stats = make_tables(weeks=7)
    up_games, up_stats = make_upcoming(games, stats, week=6)
    up_games, up_stats = make_upcoming(up_games, up_stats, week=7)
    as_of = week_kickoff(games, 6) + pd.Timedelta(hours=8)
    run = predict_upcoming(write_dataset(tmp_path, up_games, up_stats), tmp_path, as_of=as_of)
    manifest = json.loads((run / "manifest.json").read_text())
    data_age_days = (as_of - week_kickoff(games, 5)).total_seconds() / 86400
    assert data_age_days < 9
    assert manifest["counts"] == {
        "upcoming_in_dataset": 16,
        "predicted": 8,
        "skipped": 8,
        "low_history": 0,
        "past_kickoff_unresolved": 8,
    }
    assert not any("days before as-of time" in warning for warning in manifest["warnings"])
    assert any(
        "8 games are past kickoff with no recorded result" in warning
        for warning in manifest["warnings"]
    )
