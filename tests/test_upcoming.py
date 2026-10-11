"""Upcoming-game prediction tests using synthetic verified datasets; no API key needed."""

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from cfb_power.dataset_workflow import DATASET_SCHEMA, DatasetError
from cfb_power.ingestion.cfbd_canonical import PLAY_COUNT_BASIS
from cfb_power.schemas import GAME_COLUMNS, TEAM_GAME_STAT_COLUMNS
from cfb_power.upcoming import main, predict_upcoming
from cfb_power.walk_forward import walk_forward_evaluate

TEAMS = list(range(101, 117))


def make_tables(weeks=6, seed=11):
    rng = np.random.default_rng(seed)
    games, stats, game_id = [], [], 1
    for week in range(1, weeks + 1):
        order = [int(team) for team in rng.permutation(TEAMS)]
        kickoff = pd.Timestamp("2026-09-05T16:00:00Z") + pd.Timedelta(days=7 * (week - 1))
        for index in range(0, len(order), 2):
            home, away = order[index], order[index + 1]
            points = {home: int(rng.integers(10, 45)), away: int(rng.integers(10, 45))}
            games.append({
                "game_id": game_id, "season": 2026, "week": week, "game_date": kickoff.isoformat(),
                "home_team_id": home, "away_team_id": away,
                "home_score": points[home], "away_score": points[away],
                "neutral_site": False, "completed": True, "season_type": "regular",
            })
            for team, opponent, is_home in ((home, away, True), (away, home, False)):
                plays = int(rng.integers(55, 80))
                passing, rushing = int(rng.integers(120, 300)), int(rng.integers(50, 200))
                stats.append({
                    "game_id": game_id, "team_id": team, "opponent_id": opponent, "is_home": is_home,
                    "points": points[team], "points_allowed": points[opponent],
                    "total_yards": passing + rushing, "offensive_plays": plays,
                    "passing_yards": passing, "rushing_yards": rushing,
                    "turnovers": int(rng.integers(0, 4)), "takeaways": int(rng.integers(0, 4)),
                    "third_down_attempts": 12, "third_down_conversions": int(rng.integers(2, 10)),
                    "red_zone_attempts": np.nan, "red_zone_touchdowns": np.nan,
                    "penalties": int(rng.integers(2, 10)), "penalty_yards": int(rng.integers(15, 90)),
                    "time_of_possession_seconds": int(rng.integers(1500, 2100)),
                    "offensive_plays_basis": PLAY_COUNT_BASIS,
                })
            game_id += 1
    return pd.DataFrame(games, columns=[*GAME_COLUMNS, "season_type"]), pd.DataFrame(
        stats, columns=[*TEAM_GAME_STAT_COLUMNS, "offensive_plays_basis"]
    )


def make_upcoming(games, stats, week=6):
    """Turn one week into not-yet-played games: no scores, no team statistics."""
    games = games.copy()
    mask = games["week"] == week
    stats = stats[~stats["game_id"].isin(games.loc[mask, "game_id"])].copy()
    games.loc[mask, ["home_score", "away_score"]] = np.nan
    games.loc[mask, "completed"] = False
    return games, stats


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_dataset(tmp_path, games, stats, name="ds1", inputs=None):
    directory = tmp_path / "datasets" / name
    directory.mkdir(parents=True)
    outputs = []
    for filename, frame in (("games.csv", games), ("team_game_stats.csv", stats)):
        frame.to_csv(directory / filename, index=False)
        outputs.append({"file": filename, "rows": len(frame), "sha256": sha(directory / filename)})
    (directory / "manifest.json").write_text(json.dumps({
        "dataset_id": name, "schema_version": DATASET_SCHEMA, "status": "ready",
        "outputs": outputs, "inputs": inputs or [],
    }))
    return directory


def make_snapshot(tmp_path, names):
    raw = tmp_path / "raw" / "snap1"
    raw.mkdir(parents=True)
    schedule = [{"homeId": team, "homeTeam": name, "awayId": team, "awayTeam": name} for team, name in names.items()]
    (raw / "games.json").write_text(json.dumps(schedule))
    snapshot = tmp_path / "snapshots" / "snap1"
    snapshot.mkdir(parents=True)
    (snapshot / "manifest.json").write_text(json.dumps({
        "run_id": "snap1", "raw_archive": str(raw),
        "requests": [{"file": "games.json", "sha256": sha(raw / "games.json")}],
    }))
    return {"run_id": "snap1", "path": str(snapshot), "manifest_sha256": sha(snapshot / "manifest.json")}


def week_kickoff(games, week):
    return pd.Timestamp(games.loc[games["week"] == week, "game_date"].iloc[0])


def test_forecasts_match_the_walk_forward_rolling_average_exactly(tmp_path):
    games, stats = make_tables()
    full = pd.read_csv(walk_forward_evaluate(write_dataset(tmp_path, games, stats, "full"), tmp_path) / "predictions.csv")
    expected = full[full["period"] == "2026 regular week 6"].set_index("game_id")
    up_games, up_stats = make_upcoming(games, stats)
    as_of = week_kickoff(games, 6) - pd.Timedelta(days=1)
    run = predict_upcoming(write_dataset(tmp_path, up_games, up_stats, "up"), tmp_path, as_of=as_of)
    actual = pd.read_csv(run / "predictions.csv").set_index("game_id")
    assert len(actual) == 8 and set(actual.index) == set(expected.index)
    expected = expected.loc[actual.index]
    assert np.allclose(actual["home_expected_points"], expected["rolling_average_home_points"])
    assert np.allclose(actual["away_expected_points"], expected["rolling_average_away_points"])
    assert np.allclose(actual["home_win_probability"], expected["rolling_average_home_win_probability"])
    assert np.allclose(actual["projected_spread_home"], actual["home_expected_points"] - actual["away_expected_points"])
    assert np.allclose(actual["projected_total"], actual["home_expected_points"] + actual["away_expected_points"])


def test_outputs_manifest_and_workbook_are_written_without_touching_the_dataset(tmp_path):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    dataset = write_dataset(tmp_path, up_games, up_stats)
    before = {path: sha(path) for path in dataset.iterdir()}
    run = predict_upcoming(dataset, tmp_path, as_of=week_kickoff(games, 6) - pd.Timedelta(days=1))
    manifest = json.loads((run / "manifest.json").read_text())
    assert run.parent == tmp_path / "runs" / "ds1" and run.name.startswith("predictions_")
    assert manifest["model"] == "rolling_average" and manifest["counts"]["predicted"] == 8
    for item in manifest["outputs"]:
        assert sha(run / item["file"]) == item["sha256"]
    workbook = load_workbook(run / "upcoming_predictions.xlsx")
    assert workbook.sheetnames == ["README", "Predictions"]
    assert workbook["Predictions"].max_row == 9
    assert {path: sha(path) for path in dataset.iterdir()} == before


def test_games_that_have_started_are_skipped_with_a_reason(tmp_path):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    ids = up_games.loc[up_games["week"] == 6, "game_id"].tolist()
    kickoff = week_kickoff(games, 6)
    up_games.loc[up_games["game_id"].isin(ids[:4]), "game_date"] = kickoff.isoformat()
    up_games.loc[up_games["game_id"].isin(ids[4:]), "game_date"] = (kickoff + pd.Timedelta(hours=4)).isoformat()
    run = predict_upcoming(write_dataset(tmp_path, up_games, up_stats), tmp_path, as_of=kickoff + pd.Timedelta(hours=1))
    predictions = pd.read_csv(run / "predictions.csv")
    skipped = pd.read_csv(run / "skipped.csv")
    assert set(predictions["game_id"]) == set(ids[4:]) and set(skipped["game_id"]) == set(ids[:4])
    assert skipped["reason"].str.contains("kickoff has passed").all()
    assert json.loads((run / "manifest.json").read_text())["skip_reasons"] == {
        "kickoff has passed or game is in progress": 4
    }


def test_low_history_is_flagged_and_threshold_is_configurable(tmp_path):
    games, stats = make_tables(weeks=3)
    up_games, up_stats = make_upcoming(games, stats, week=3)
    as_of = week_kickoff(games, 3) - pd.Timedelta(days=1)
    dataset = write_dataset(tmp_path, up_games, up_stats)
    flagged = pd.read_csv(predict_upcoming(dataset, tmp_path, as_of=as_of) / "predictions.csv")
    assert flagged["low_history"].all() and set(flagged["home_games_played"]) == {2}
    relaxed = pd.read_csv(predict_upcoming(dataset, tmp_path, as_of=as_of, min_history_games=2) / "predictions.csv")
    assert not relaxed["low_history"].any()


def test_team_without_history_falls_back_to_league_means(tmp_path):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    extra = up_games[up_games["week"] == 6].iloc[[0]].copy()
    extra[["game_id", "home_team_id", "away_team_id"]] = [999, 998, 997]
    up_games = pd.concat([up_games, extra], ignore_index=True)
    run = predict_upcoming(write_dataset(tmp_path, up_games, up_stats), tmp_path, as_of=week_kickoff(games, 6) - pd.Timedelta(days=1))
    row = pd.read_csv(run / "predictions.csv").set_index("game_id").loc[999]
    completed = games[games["week"] <= 5]
    assert row["home_expected_points"] == pytest.approx(completed["home_score"].mean())
    assert row["away_expected_points"] == pytest.approx(completed["away_score"].mean())
    assert row["low_history"] and row["home_games_played"] == 0
    assert 0 < row["home_win_probability"] < 1


def test_team_names_come_from_verified_raw_archives_only(tmp_path):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    snapshot = make_snapshot(tmp_path, {team: f"Team {team}" for team in TEAMS})
    as_of = week_kickoff(games, 6) - pd.Timedelta(days=1)
    run = predict_upcoming(write_dataset(tmp_path, up_games, up_stats, "named", [snapshot]), tmp_path, as_of=as_of)
    named = pd.read_csv(run / "predictions.csv")
    assert (named["home_team"] == "Team " + named["home_team_id"].astype(str)).all()
    Path(snapshot["path"], "manifest.json").write_text("{}")
    run = predict_upcoming(write_dataset(tmp_path, up_games, up_stats, "tampered", [snapshot]), tmp_path, as_of=as_of)
    unnamed = pd.read_csv(run / "predictions.csv")
    manifest = json.loads((run / "manifest.json").read_text())
    assert unnamed["home_team"].isna().all() and len(unnamed) == 8
    assert any("team names unavailable" in warning for warning in manifest["warnings"])


def test_stale_data_produces_a_warning(tmp_path):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    later = (week_kickoff(games, 6) + pd.Timedelta(days=30)).isoformat()
    up_games.loc[up_games["week"] == 6, "game_date"] = later
    as_of = week_kickoff(games, 5) + pd.Timedelta(days=20)
    run = predict_upcoming(write_dataset(tmp_path, up_games, up_stats), tmp_path, as_of=as_of)
    assert any("days before as-of time" in warning for warning in json.loads((run / "manifest.json").read_text())["warnings"])


def test_nothing_to_predict_fails_without_creating_a_run(tmp_path):
    games, stats = make_tables()
    with pytest.raises(DatasetError, match="no upcoming games"):
        predict_upcoming(write_dataset(tmp_path, games, stats, "none"), tmp_path)
    up_games, up_stats = make_upcoming(games, stats)
    with pytest.raises(DatasetError, match="kickoff has passed"):
        predict_upcoming(write_dataset(tmp_path, up_games, up_stats, "past"), tmp_path, as_of=week_kickoff(games, 6) + pd.Timedelta(days=1))
    assert not (tmp_path / "runs").exists()


@pytest.mark.parametrize("bad", ["2026-10-10 12:00", "not a time"])
def test_as_of_must_be_a_timestamp_with_timezone(tmp_path, bad):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    with pytest.raises(ValueError):
        predict_upcoming(write_dataset(tmp_path, up_games, up_stats), tmp_path, as_of=bad)
    with pytest.raises(ValueError):
        predict_upcoming(tmp_path / "x", tmp_path, min_history_games=-1)


def test_tampered_dataset_is_rejected(tmp_path):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    dataset = write_dataset(tmp_path, up_games, up_stats)
    with (dataset / "games.csv").open("a", encoding="utf-8") as handle:
        handle.write("\n")
    with pytest.raises(DatasetError, match="hash mismatch"):
        predict_upcoming(dataset, tmp_path, as_of=week_kickoff(games, 6) - pd.Timedelta(days=1))


def test_cli_prints_forecasts_and_reports_failures(tmp_path, capsys):
    games, stats = make_tables()
    up_games, up_stats = make_upcoming(games, stats)
    dataset = write_dataset(tmp_path, up_games, up_stats)
    as_of = (week_kickoff(games, 6) - pd.Timedelta(days=1)).isoformat()
    assert main(["--dataset", str(dataset), "--data-dir", str(tmp_path), "--as-of", as_of]) == 0
    output = capsys.readouterr().out
    assert "Predicted 8 games; skipped 0" in output and "WARNING" not in output
    assert main(["--dataset", str(dataset), "--data-dir", str(tmp_path), "--as-of", "bad"]) == 1
    assert main(["--dataset", str(tmp_path / "missing"), "--data-dir", str(tmp_path), "--as-of", as_of]) == 1
    assert "Upcoming predictions failed" in capsys.readouterr().err
