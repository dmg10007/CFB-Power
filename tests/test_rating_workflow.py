import hashlib
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from cfb_power.rating_workflow import (
    STAT_FIELDS, audit_source_coverage, load_rating_dataset, main, run_rating_workflow,
)

CUTOFF = "2026-10-19T12:00:00Z"
BASIS = "passing_attempts_plus_rushing_attempts_proxy"


@pytest.fixture
def dataset(tmp_path):
    source = tmp_path / "dataset"
    source.mkdir()
    games = pd.DataFrame([{"game_id":1,"season":2026,"week":1,
        "game_date":"2026-10-01T12:00:00Z","home_team_id":1,"away_team_id":2,
        "home_score":35,"away_score":21,"neutral_site":True,"completed":True}])
    rows = []
    for team,opponent,points,allowed,home in [(1,2,35,21,True),(2,1,21,35,False)]:
        row = dict.fromkeys(sorted(STAT_FIELDS), 0)
        row.update(game_id=1,team_id=team,opponent_id=opponent,points=points,
                   points_allowed=allowed,is_home=home,total_yards=350,offensive_plays=60,
                   passing_yards=200,rushing_yards=150,third_down_attempts=10,
                   third_down_conversions=4,time_of_possession_seconds=1800,
                   red_zone_attempts=None,red_zone_touchdowns=None,offensive_plays_basis=BASIS)
        rows.append(row)
    games.to_csv(source/"games.csv",index=False)
    pd.DataFrame(rows).to_csv(source/"team_game_stats.csv",index=False)
    manifest = {"status":"ready","schema_version":"dataset-basic-v1","dataset_id":"fixture",
                "created_at_utc":"2026-10-19T10:00:00Z","offensive_plays_basis":BASIS,
                "outputs":[]}
    for name in ("games.csv","team_game_stats.csv"):
        manifest["outputs"].append({"file":name,"rows":len(pd.read_csv(source/name)),
            "sha256":hashlib.sha256((source/name).read_bytes()).hexdigest()})
    (source/"manifest.json").write_text(json.dumps(manifest))
    return source


def rewrite_manifest(source, change):
    path = source/"manifest.json"
    manifest = json.loads(path.read_text())
    change(manifest)
    path.write_text(json.dumps(manifest))


def rehash(source, name):
    def update(manifest):
        for entry in manifest["outputs"]:
            if entry["file"]==name:
                entry["sha256"]=hashlib.sha256((source/name).read_bytes()).hexdigest()
                entry["rows"]=len(pd.read_csv(source/name))
    rewrite_manifest(source,update)


def test_verified_dataset_loads(dataset):
    manifest,games,stats,digest=load_rating_dataset(dataset)
    assert len(games)==1 and len(stats)==2
    assert len(digest)==64 and manifest["dataset_id"]=="fixture"


@pytest.mark.parametrize("mutation",["hash","rows","schema","basis","outputs"])
def test_bad_manifest_rejected(dataset,mutation):
    def change(manifest):
        if mutation=="hash": manifest["outputs"][0]["sha256"]="0"*64
        if mutation=="rows": manifest["outputs"][0]["rows"]=10
        if mutation=="schema": manifest["schema_version"]="unknown"
        if mutation=="basis": manifest["offensive_plays_basis"]="real_snaps"
        if mutation=="outputs": manifest["outputs"].append(manifest["outputs"][0])
    rewrite_manifest(dataset,change)
    with pytest.raises(ValueError): load_rating_dataset(dataset)


@pytest.mark.parametrize("mutation",["opponent","score","team","duplicate","unfinished"])
def test_inconsistent_stats_rejected(dataset,mutation):
    path=dataset/"team_game_stats.csv"
    frame=pd.read_csv(path)
    if mutation=="opponent": frame.loc[0,"opponent_id"]=99
    if mutation=="score": frame.loc[0,"points"]=999
    if mutation=="team": frame.loc[0,"team_id"]=99
    if mutation=="duplicate": frame=pd.concat([frame,frame.iloc[[0]]],ignore_index=True)
    if mutation=="unfinished": frame.loc[0,"game_id"]=99
    frame.to_csv(path,index=False)
    rehash(dataset,path.name)
    with pytest.raises(ValueError): load_rating_dataset(dataset)


def test_build_publishes_workbook_and_hashes_without_touching_inputs(dataset,tmp_path):
    before={p.name:p.read_bytes() for p in dataset.iterdir()}
    output=run_rating_workflow(dataset,tmp_path/"data",season=2026,as_of=CUTOFF)
    manifest=json.loads((output/"manifest.json").read_text())
    assert manifest["experimental"] and manifest["external_api_requests"]==0
    assert manifest["point_in_time_provenance_verified"] is False
    for item in manifest["outputs"]:
        assert hashlib.sha256((output/item["file"]).read_bytes()).hexdigest()==item["sha256"]
    assert {p.name:p.read_bytes() for p in dataset.iterdir()}==before
    workbook=load_workbook(output/"power_ratings.xlsx",read_only=True)
    assert workbook.sheetnames==["Team Ratings","Coverage"]
    workbook.close()
    assert len(pd.read_csv(output/"team_ratings.csv"))==2


def test_two_builds_preserve_old_snapshot(dataset,tmp_path):
    a=run_rating_workflow(dataset,tmp_path/"data",season=2026,as_of=CUTOFF)
    original={p.name:p.read_bytes() for p in a.iterdir()}
    b=run_rating_workflow(dataset,tmp_path/"data",season=2026,as_of=CUTOFF)
    assert a!=b and a.exists() and b.exists()
    assert original=={p.name:p.read_bytes() for p in a.iterdir()}


def test_audit_only_does_not_generate_rankings(dataset,tmp_path):
    output=run_rating_workflow(dataset,tmp_path/"data",season=2026,as_of=CUTOFF,build=False)
    assert not (output/"team_ratings.csv").exists()
    assert json.loads((output/"manifest.json").read_text())["rating_model"] is None


def test_proxy_and_red_zone_are_labeled(dataset):
    _,games,stats,_=load_rating_dataset(dataset)
    audit=audit_source_coverage(games,stats,season=2026,as_of=CUTOFF)
    features=pd.DataFrame(audit["features"]).set_index("feature")
    assert features.loc["yards_per_offensive_play_proxy","status"]=="proxy"
    assert features.loc["red_zone_td_rate","status"]=="missing"
    assert features.loc["third_down_conversion_rate","status"]=="available"


def test_duplicate_csv_headers_rejected(dataset):
    path=dataset/"team_game_stats.csv"
    text=path.read_text().splitlines()
    text[0]=text[0].replace("points_allowed","points",1)
    path.write_text("\n".join(text)+"\n")
    rehash(dataset,path.name)
    with pytest.raises(ValueError,match="Duplicate CSV"):
        load_rating_dataset(dataset)


def test_cli_audit_runs_offline(dataset,tmp_path,monkeypatch):
    import socket
    def forbidden(*args,**kwargs):
        raise AssertionError("Network access is forbidden")
    monkeypatch.setattr(socket,"create_connection",forbidden)
    assert main(["audit","--dataset",str(dataset),"--data-dir",str(tmp_path/"data"),
                 "--season","2026","--as-of",CUTOFF])==0


def test_failed_publication_cleans_staging(dataset,tmp_path,monkeypatch):
    import cfb_power.rating_workflow as workflow
    def fail(*args,**kwargs):
        raise OSError("simulated workbook failure")
    monkeypatch.setattr(workflow,"_write_workbook",fail)
    with pytest.raises(OSError):
        run_rating_workflow(dataset,tmp_path/"data",season=2026,as_of=CUTOFF)
    assert list((tmp_path/"data/rating_snapshots").iterdir())==[]


def test_input_mutation_before_publication_fails(dataset,tmp_path,monkeypatch):
    import cfb_power.rating_workflow as workflow
    original=workflow._write_workbook
    def mutate(*args,**kwargs):
        original(*args,**kwargs)
        with (dataset/"games.csv").open("a") as handle:
            handle.write("\n")
    monkeypatch.setattr(workflow,"_write_workbook",mutate)
    with pytest.raises(ValueError,match="changed"):
        run_rating_workflow(dataset,tmp_path/"data",season=2026,as_of=CUTOFF)
    assert list((tmp_path/"data/rating_snapshots").iterdir())==[]


def test_cli_invalid_input_returns_nonzero(tmp_path):
    assert main(["build","--dataset",str(tmp_path/"absent"),"--season","2026","--as-of",CUTOFF])==1
